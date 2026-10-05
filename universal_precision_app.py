import io
import re
import zipfile
from collections import defaultdict
from copy import copy

import fitz
import streamlit as st
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Universal Precision Validation", page_icon="🧪", layout="wide")
st.title("🧪 Universal Precision Validation Automation")
st.caption("Template-flexible precision population | INTRA + dynamic INTER runs")

NUM = r'[<>]?\s*-?\d+(?:\.\d+)?(?:[Ee][+-]?\d+)?'
RECORD_RE = re.compile(
    rf'\bPR\s+(INTRA|INTER)\s+([A-Za-z0-9_.-]+)\s+({NUM})\s*([^\s]+)',
    re.I
)

def normalize_analyte(name):
    name = re.sub(r'[^A-Za-z0-9]+', '_', str(name).strip())
    return re.sub(r'_+', '_', name).strip('_') or "Analyte"

def level_from_sid(mode, sid):
    s = str(sid).upper().strip()
    if mode == "INTRA":
        m = re.match(r'L([123])(?:[.]?01)?$', s) or re.search(r'L([123])', s)
        return int(m.group(1)) if m else None
    if re.fullmatch(r'[34]0[1-3]', s):
        return int(s[-1])
    m = re.fullmatch(r'R20([1-3])', s)
    if m:
        return int(m.group(1))
    m = re.search(r'0([1-3])$', s)
    return int(m.group(1)) if m else None

def extract_pdf_records(uploaded_file, expected_mode):
    doc = fitz.open(stream=uploaded_file.getvalue(), filetype="pdf")
    records = []
    for page_no, page in enumerate(doc, 1):
        text = page.get_text("text") or ""
        if not text:
            continue

        assay = None
        m_assay = re.search(r'Tests:\s*([^;\n]+)', text, re.I)
        if m_assay:
            assay = m_assay.group(1).strip()

        if not assay:
            lines = [x.strip() for x in text.splitlines() if x.strip()]
            for i, line in enumerate(lines):
                if line.lower().startswith("test status") and i + 1 < len(lines):
                    if lines[i + 1].lower() != "assay report":
                        assay = lines[i + 1]
                        break

        for m in RECORD_RE.finditer(text):
            mode = m.group(1).upper()
            if mode != expected_mode:
                continue
            sid = m.group(2)
            raw = m.group(3).replace(" ", "")
            try:
                result = float(raw.replace("<", "").replace(">", ""))
            except ValueError:
                continue
            level = level_from_sid(mode, sid)
            if level is None:
                continue
            records.append({
                "analyte": assay or "Unknown",
                "mode": mode,
                "sid": sid,
                "level": level,
                "result": result,
                "unit": m.group(4),
                "page": page_no
            })
    doc.close()
    return records

def group_records(records):
    grouped = defaultdict(lambda: defaultdict(list))
    units = {}
    for r in records:
        grouped[r["analyte"]][r["level"]].append(r["result"])
        units.setdefault(r["analyte"], r["unit"])
    return grouped, units

def merged_safe_cell(ws, row, col):
    cell = ws.cell(row, col)
    if not isinstance(cell, MergedCell):
        return cell
    for rng in ws.merged_cells.ranges:
        if cell.coordinate in rng:
            min_col, min_row, max_col, max_row = rng.bounds
            if row == min_row and col == min_col:
                return ws.cell(min_row, min_col)
            raise ValueError(f"Merged range {rng} blocks cell {cell.coordinate}.")
    return cell

def put(ws, row, col, value):
    merged_safe_cell(ws, row, col).value = value

def copy_row_style(ws, source_row, target_row):
    for col in range(1, ws.max_column + 1):
        src = ws.cell(source_row, col)
        dst = ws.cell(target_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        dst.number_format = src.number_format
        dst.alignment = copy(src.alignment)
        dst.protection = copy(src.protection)
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height

def text_at(ws, row, col):
    return str(ws.cell(row, col).value or "").strip().lower()

def find_column_by_labels(ws, labels, max_scan_rows=60):
    labels = {x.lower() for x in labels}
    for r in range(1, min(ws.max_row, max_scan_rows) + 1):
        for c in range(1, min(ws.max_column, 40) + 1):
            if text_at(ws, r, c) in labels:
                return c
    return None

def find_stat_rows(ws):
    result = {}
    for r in range(1, min(ws.max_row, 120) + 1):
        row_text = " ".join(text_at(ws, r, c) for c in range(1, min(ws.max_column, 30) + 1))
        if re.search(r'\bmean\b', row_text) and "mean" not in result:
            result["mean"] = r
        if re.search(r'\b(?:sd|s\.d\.?)\b', row_text) and "sd" not in result:
            result["sd"] = r
        if re.search(r'\bcv\s*%', row_text) and "cv" not in result:
            result["cv"] = r
        if re.search(r'\b(?:n|count)\b', row_text) and "n" not in result:
            result["n"] = r
    return result

def detect_layout(ws):
    # The universal detector first searches labels. If a template uses a different
    # visual layout but keeps the same population concept, it falls back to the
    # proven standard precision layout.
    replicate_col = find_column_by_labels(ws, {"replicate", "rep", "replication"})
    stats = find_stat_rows(ws)

    layout = {
        "inter_cols": [2, 3, 4],
        "intra_cols": [7, 8, 9],
        "replicate_col": replicate_col or 6,
        "start_row": 7,
        "end_row": 26,
        "n_row": stats.get("n", 28),
        "mean_row": stats.get("mean", 29),
        "sd_row": stats.get("sd", 30),
        "cv_row": stats.get("cv", 31),
    }

    # Detect likely data header row from Replicate. Keep the 20-row convention
    # immediately below it, which is common to the precision templates.
    if replicate_col:
        for r in range(1, min(ws.max_row, 60) + 1):
            if text_at(ws, r, replicate_col) in {"replicate", "rep", "replication"}:
                layout["start_row"] = r + 1
                layout["end_row"] = r + 20
                break

    return layout

def ensure_rows(ws, layout, required_rows):
    current_rows = layout["end_row"] - layout["start_row"] + 1
    extra = max(0, required_rows - current_rows)
    if not extra:
        return layout

    stats_anchor = min(layout["n_row"], layout["mean_row"], layout["sd_row"], layout["cv_row"])
    source_row = layout["end_row"]
    ws.insert_rows(stats_anchor, amount=extra)

    for r in range(source_row + 1, source_row + extra + 1):
        copy_row_style(ws, source_row, r)

    for key in ("n_row", "mean_row", "sd_row", "cv_row"):
        if layout[key] >= stats_anchor:
            layout[key] += extra
    layout["end_row"] += extra
    return layout

def clear_data_area(ws, layout):
    for row in range(layout["start_row"], layout["end_row"] + 1):
        for col in layout["inter_cols"] + layout["intra_cols"]:
            put(ws, row, col, None)
        put(ws, row, layout["replicate_col"], row - layout["start_row"] + 1)

def set_formula(ws, row, col, formula):
    put(ws, row, col, formula)

def populate_workbook(template_bytes, analyte, unit, n_levels, intra_by_level, inter_runs_by_level):
    wb = load_workbook(io.BytesIO(template_bytes))
    ws = wb.active
    layout = detect_layout(ws)

    # Each level has Run 1 (5 INTRA values) plus one 5-value block for every
    # uploaded INTER PDF. The maximum number of rows is therefore determined
    # from the actual uploaded PDFs, not from a fixed template limit.
    inter_pdf_count = max(
        [len(inter_runs_by_level.get(level, [])) for level in range(1, n_levels + 1)] or [0]
    )
    required_inter_rows = 5 * (1 + inter_pdf_count)
    required_rows = max(20, required_inter_rows)

    layout = ensure_rows(ws, layout, required_rows)
    clear_data_area(ws, layout)

    # Keep common template title/unit locations when present.
    try:
        if ws["A3"].value is not None:
            put(ws, 3, 1, f"Precision Studies of - {analyte}")
    except Exception:
        pass
    if unit:
        # Existing precision templates commonly use E7 for unit.
        try:
            if ws.max_column >= 5:
                put(ws, layout["start_row"], 5, unit)
        except Exception:
            pass

    # INTRA: 20 observations per QC level.
    for level in range(1, n_levels + 1):
        values = list(intra_by_level.get(level, []))[:20]
        col = layout["intra_cols"][level - 1]
        for i, value in enumerate(values):
            put(ws, layout["start_row"] + i, col, value)

    # INTER: Run 1 = first 5 INTRA; Run 2+ = each uploaded INTER PDF.
    for level in range(1, n_levels + 1):
        col = layout["inter_cols"][level - 1]
        values = list(intra_by_level.get(level, []))[:5]
        for run_values in inter_runs_by_level.get(level, []):
            values.extend(list(run_values)[:5])
        for i, value in enumerate(values):
            put(ws, layout["start_row"] + i, col, value)

    s, e = layout["start_row"], layout["end_row"]
    for col in layout["inter_cols"] + layout["intra_cols"]:
        letter = get_column_letter(col)
        set_formula(ws, layout["n_row"], col, f'=COUNT({letter}{s}:{letter}{e})')
        set_formula(ws, layout["mean_row"], col, f'=IFERROR(AVERAGE({letter}{s}:{letter}{e}),"")')
        set_formula(ws, layout["sd_row"], col, f'=IFERROR(STDEV({letter}{s}:{letter}{e}),"")')
        set_formula(
            ws,
            layout["cv_row"],
            col,
            f'=IFERROR(({letter}{layout["sd_row"]}/{letter}{layout["mean_row"]})*100,"")'
        )

    if n_levels == 2:
        for col in (layout["inter_cols"][2], layout["intra_cols"][2]):
            for row in range(layout["start_row"], layout["cv_row"] + 1):
                put(ws, row, col, None)

    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
    except Exception:
        pass
    return wb

def wb_bytes(wb):
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

st.sidebar.header("Validation settings")
level_mode = st.sidebar.selectbox("QC level configuration", ["Auto Detect", "2 Levels", "3 Levels"])

st.info(
    "Universal mapping: INTRA = 20 replicates/level. "
    "INTER Run 1 = first 5 INTRA. Run 2+ = one run for each uploaded INTER PDF, "
    "in upload order."
)

template = st.file_uploader(
    "1. Upload Precision Excel template", type=["xlsx", "xlsm"]
)
intra_files = st.file_uploader(
    "2. Upload INTRA PDF(s)", type=["pdf"], accept_multiple_files=True
)
inter_files = st.file_uploader(
    "3. Upload INTER PDF(s) in RUN order (optional)",
    type=["pdf"], accept_multiple_files=True
)

if st.button("🔍 Process Precision Study", type="primary", use_container_width=True):
    if not template or not intra_files:
        st.error("Please upload the Excel template and at least one INTRA PDF.")
        st.stop()

    with st.spinner("Reading PDF result values..."):
        intra_records = []
        for f in intra_files:
            intra_records.extend(extract_pdf_records(f, "INTRA"))

        intra_group, units = group_records(intra_records)

        inter_file_groups = []
        for f in inter_files or []:
            records = extract_pdf_records(f, "INTER")
            grouped, grouped_units = group_records(records)
            inter_file_groups.append((f.name, grouped, grouped_units))

    if not intra_records:
        st.error("No INTRA Result records detected. Check the PDF format.")
        st.stop()

    detected_levels = sorted(
        {lvl for a in intra_group.values() for lvl in a if lvl in (1, 2, 3)}
    )
    n_levels = (
        2 if level_mode == "2 Levels"
        else 3 if level_mode == "3 Levels"
        else (3 if 3 in detected_levels else 2)
    )

    st.subheader("Extraction summary")
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Analytes", len(intra_group))
    c2.metric("QC levels", n_levels)
    c3.metric("INTRA records", len(intra_records))
    c4.metric("INTER PDFs", len(inter_file_groups))

    summary = []
    for analyte in sorted(intra_group):
        for level in range(1, n_levels + 1):
            counts = [
                len(g.get(analyte, {}).get(level, []))
                for _, g, _ in inter_file_groups
            ]
            summary.append({
                "Analyte": analyte,
                "Level": f"L{level}",
                "INTRA": len(intra_group[analyte].get(level, [])),
                "INTER PDF counts": ", ".join(map(str, counts)) if counts else "-",
                "Population": "Run 1 = INTRA first 5; Run 2+ = each INTER PDF"
            })
    st.dataframe(summary, use_container_width=True)

    outputs = {}
    warnings = []

    for analyte in sorted(intra_group):
        intra_by_level = {
            level: list(intra_group[analyte].get(level, []))
            for level in range(1, n_levels + 1)
        }

        for level in range(1, n_levels + 1):
            if len(intra_by_level[level]) < 20:
                st.error(
                    f"{analyte} L{level}: only "
                    f"{len(intra_by_level[level])}/20 INTRA results detected."
                )
                st.stop()

        inter_runs_by_level = defaultdict(list)

        for filename, grouped, _ in inter_file_groups:
            for level in range(1, n_levels + 1):
                vals = list(grouped.get(analyte, {}).get(level, []))[:5]
                if len(vals) == 5:
                    inter_runs_by_level[level].append(vals)
                elif vals:
                    warnings.append(
                        f"{analyte} L{level}: {filename} has "
                        f"{len(vals)}/5 INTER values; this run was not populated."
                    )
                else:
                    warnings.append(
                        f"{analyte} L{level}: {filename} has no INTER values; "
                        f"this run was not populated."
                    )

        wb = populate_workbook(
            template.getvalue(),
            analyte,
            units.get(analyte, ""),
            n_levels,
            intra_by_level,
            inter_runs_by_level
        )

        outputs[f"{normalize_analyte(analyte)}_Precision.xlsx"] = wb_bytes(wb)

    zip_b = io.BytesIO()
    with zipfile.ZipFile(zip_b, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in outputs.items():
            z.writestr(name, data)
    zip_b.seek(0)

    st.success(f"Created {len(outputs)} analyte workbook(s).")

    if warnings:
        st.warning("Review these items:")
        for warning in warnings:
            st.write("• " + warning)

    st.download_button(
        "⬇️ DOWNLOAD ALL ANALYTES — ZIP",
        data=zip_b.getvalue(),
        file_name="Universal_Precision_Validation_Output.zip",
        mime="application/zip",
        use_container_width=True
    )
